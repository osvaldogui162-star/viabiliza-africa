from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

from app.application.interfaces.excel_parser import ExcelCostItemRow, IExcelParser
from app.domain.exceptions.domain_exceptions import ValidationError


class OpenpyxlExcelParser(IExcelParser):
    """
    Parser Excel para importação OPEX/CAPEX (UC13).
    Colunas esperadas: tipo, categoria, descricao, quantidade, unidade, preco_unitario, fornecedor
    """

    HEADERS = {
        "tipo": "item_type",
        "type": "item_type",
        "categoria": "category",
        "category": "category",
        "descricao": "description",
        "descrição": "description",
        "description": "description",
        "quantidade": "quantity",
        "quantity": "quantity",
        "unidade": "unit",
        "unit": "unit",
        "preco_unitario": "unit_price",
        "preço_unitário": "unit_price",
        "unit_price": "unit_price",
        "fornecedor": "supplier_name",
        "supplier": "supplier_name",
        "nif": "supplier_nif",
        "nif_fornecedor": "supplier_nif",
        "supplier_nif": "supplier_nif",
    }

    def parse_cost_items(self, file_bytes: bytes) -> list[ExcelCostItemRow]:
        try:
            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        except Exception as exc:
            raise ValidationError(f"Ficheiro Excel inválido: {exc}") from exc

        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            raise ValidationError("Excel deve ter cabeçalho e pelo menos uma linha de dados")

        header_map = self._map_headers(rows[0])
        data_rows = rows[1:]
        if len(data_rows) > self.MAX_ROWS:
            raise ValidationError(
                f"Importação limitada a {self.MAX_ROWS} linhas por upload"
            )

        items: list[ExcelCostItemRow] = []
        for index, row in enumerate(data_rows, start=2):
            if self._is_empty_row(row):
                continue
            try:
                items.append(self._parse_row(row, header_map, index))
            except ValidationError:
                raise
            except Exception as exc:
                raise ValidationError(f"Erro na linha {index}: {exc}") from exc

        if not items:
            raise ValidationError("Nenhum item válido encontrado no Excel")
        return items

    def _map_headers(self, header_row: tuple) -> dict[int, str]:
        mapping: dict[int, str] = {}
        for col_index, cell in enumerate(header_row):
            if cell is None:
                continue
            key = str(cell).strip().lower()
            if key in self.HEADERS:
                mapping[col_index] = self.HEADERS[key]
        required = {"item_type", "category", "description", "quantity", "unit_price"}
        found = set(mapping.values())
        if not required.issubset(found):
            missing = required - found
            raise ValidationError(
                f"Colunas obrigatórias em falta no Excel: {', '.join(sorted(missing))}"
            )
        return mapping

    def _parse_row(
        self, row: tuple, header_map: dict[int, str], line_number: int
    ) -> ExcelCostItemRow:
        data: dict = {}
        for col_index, field_name in header_map.items():
            value = row[col_index] if col_index < len(row) else None
            data[field_name] = value

        item_type = str(data.get("item_type", "")).strip().lower()
        if item_type not in ("opex", "capex"):
            raise ValidationError(f"Linha {line_number}: tipo deve ser 'opex' ou 'capex'")

        quantity = self._to_decimal(data.get("quantity"), line_number, "quantidade")
        unit_price = self._to_decimal(data.get("unit_price"), line_number, "preço unitário")
        if quantity <= 0:
            raise ValidationError(f"Linha {line_number}: quantidade deve ser > 0")
        if unit_price < 0:
            raise ValidationError(f"Linha {line_number}: preço unitário inválido")

        return ExcelCostItemRow(
            item_type=item_type,
            category=str(data.get("category", "")).strip(),
            description=str(data.get("description", "")).strip(),
            quantity=quantity,
            unit=str(data.get("unit") or "un").strip(),
            unit_price=unit_price,
            supplier_name=(
                str(data["supplier_name"]).strip() if data.get("supplier_name") else None
            ),
            supplier_nif=(
                str(data["supplier_nif"]).strip().upper().replace(" ", "")
                if data.get("supplier_nif")
                else None
            ),
        )

    @staticmethod
    def _to_decimal(value, line_number: int, field: str) -> Decimal:
        if value is None or str(value).strip() == "":
            raise ValidationError(f"Linha {line_number}: {field} é obrigatório")
        try:
            return Decimal(str(value))
        except (InvalidOperation, ValueError) as exc:
            raise ValidationError(f"Linha {line_number}: {field} inválido") from exc

    @staticmethod
    def _is_empty_row(row: tuple) -> bool:
        return all(cell is None or str(cell).strip() == "" for cell in row)
